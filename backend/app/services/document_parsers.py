"""Text extraction from uploaded files, keeping where each piece of text came from.

Pure functions over bytes, so they are easy to test and can run in a worker
thread. Images, PDF pages without a text layer (scanned pages), and the pictures
on pages and slides (diagrams, charts, code, tables) are read by the AI model
instead; see DocumentIngestionService and figure_reading.py.
"""

import hashlib
import io
import math
import re
from collections import Counter
from collections.abc import Iterator
from dataclasses import dataclass, field

PDF = "application/pdf"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TXT = "text/plain"
IMAGE_TYPES = {"image/png", "image/jpeg", "image/webp"}

# A PDF page with less text than this is treated as having no text layer.
_MIN_PAGE_CHARACTERS = 25

# A picture is worth reading when it may carry course content: it covers at least this share of
# its page; or it's a strip (an equation, a line of code, a timeline) at least this long relative
# to its height; or it's the main thing on a page with almost no text. Icons, emoji, bullets and
# logos are smaller than that, and a picture on most pages is the slide template (a logo, a
# background), never content.
_PICTURE_MIN_AREA = 0.05
_STRIP_RATIO = 3.0
_STRIP_MIN_AREA = 0.01
_SPARSE_PAGE_WORDS = 15
_SPARSE_MIN_AREA = 0.015
_TEMPLATE_MIN_PAGES = 3
# Title and closing slides aren't worth reading, whatever pictures they have.
_TITLE_WORDS = 30
_CLOSING_WORDS = 12
_CLOSING = re.compile(
    r"(?i)\b(?:thank(?:s| you)|any questions|the end|end of (?:the )?(?:lecture|presentation|session|chapter))\b"
    r"|\bquestions\s*\?|\bq\s*&\s*a\b"
)


class DocumentParseError(Exception):
    """A file couldn't be read. The message is safe to show the professor."""


@dataclass
class TextUnit:
    """A natural piece of a document: a page, a slide, or a section."""

    text: str
    page_start: int | None = None
    page_end: int | None = None
    # "Slide", "Page" or "Section", for building location labels.
    kind: str | None = None
    # A section heading, for Word documents.
    heading: str | None = None


@dataclass
class Picture:
    """A picture drawn on a page or slide: which one it is, and how much of the page it covers."""

    # The same picture on several pages has the same key (so a slide template is recognised).
    key: object
    # The share of the page or slide it covers (0–1), and its longer side over its shorter one.
    area: float
    ratio: float
    # PowerPoint: the picture itself (content type, bytes), to show the AI. A PDF page is shown whole.
    image: tuple[str, bytes] | None = None


@dataclass
class FigureSpot:
    """A page or slide whose pictures may carry course content, for the AI to read (see figure_reading.py)."""

    number: int
    # "Page" or "Slide".
    kind: str
    # The page's own text (it may be too short to keep on its own, such as a title over a diagram).
    text: str
    # The pictures on it that may carry content.
    pictures: list[Picture]


@dataclass
class ParsedDocument:
    units: list[TextUnit]
    page_count: int | None = None
    # PDF pages (1-based) that have no text layer and should be transcribed.
    pages_without_text: list[int] = field(default_factory=list)
    # Pages or slides whose pictures may carry content, in order. Each has a "not in the
    # extracted text" marker in its text until the pictures are read.
    figures: list[FigureSpot] = field(default_factory=list)
    # PDF: every page's own text layer, by page number, short or empty ones included.
    page_texts: dict[int, str] = field(default_factory=dict)

    @property
    def character_count(self) -> int:
        return sum(len(unit.text) for unit in self.units)


def parse_document(data: bytes, mime_type: str) -> ParsedDocument:
    if mime_type == PDF:
        return parse_pdf(data)
    if mime_type == PPTX:
        return parse_pptx(data)
    if mime_type == DOCX:
        return parse_docx(data)
    if mime_type == TXT:
        return parse_txt(data)
    raise DocumentParseError("This file type can't be read.")


def normalize_text(text: str) -> str:
    """Tidies extracted text: consistent line breaks, no runs of blank space."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_pdf(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            # Some PDFs are "encrypted" with an empty password, only to restrict printing.
            if not reader.decrypt(""):
                raise DocumentParseError("This PDF is password-protected. Upload a copy without a password.")
        pages = list(reader.pages)
    except DocumentParseError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, OSError) as error:
        raise DocumentParseError("This PDF couldn't be opened. It may be damaged.") from error

    units: list[TextUnit] = []
    without_text: list[int] = []
    figures: list[FigureSpot] = []
    page_texts: dict[int, str] = {}
    keys: dict[int, object] = {}
    pictures = [_pdf_pictures(page, number, keys) for number, page in enumerate(pages, start=1)]
    template = _template_keys(pictures)
    seen: set[object] = set()
    for number, page in enumerate(pages, start=1):
        try:
            text = normalize_text(page.extract_text() or "")
        except Exception:  # a single unreadable page shouldn't sink the document
            text = ""
        page_texts[number] = text
        worth = _worth_reading(number, text, len(pages))
        content = _content_pictures(pictures[number - 1], text, template, seen) if worth else []
        if content:
            figures.append(FigureSpot(number=number, kind="Page", text=text, pictures=content))
        if len(text) < _MIN_PAGE_CHARACTERS:
            without_text.append(number)
            continue
        if content:
            # Pictures aren't text: say so, so their content is never assumed to be covered (until read).
            text = f"{text}\n{figure_marker('page', number, len(content))}"
        units.append(TextUnit(text=text, page_start=number, page_end=number, kind="Page"))
    return ParsedDocument(
        units=units, page_count=len(pages), pages_without_text=without_text, figures=figures, page_texts=page_texts
    )


def _pdf_pictures(page, number: int, keys: dict[int, object]) -> list[Picture]:
    """The pictures drawn on a PDF page, with the share of the page each covers (nothing is decoded).

    Follows the page's drawing instructions, so a picture's size is the size it's drawn at
    (inside reusable drawings too), not the size of its file. `keys` remembers each picture's
    key across pages.
    """
    from pypdf.generic import ContentStream

    box = page.mediabox
    page_area = abs(float(box.width) * float(box.height)) or 1.0
    found: list[Picture] = []

    def add(key: object, ctm: tuple[float, ...]) -> None:
        a, b, c, d = ctm[:4]
        width, height = math.hypot(a, b), math.hypot(c, d)
        ratio = max(width, height) / max(min(width, height), 1e-9)
        found.append(Picture(key=key, area=min(1.0, abs(a * d - b * c) / page_area), ratio=ratio))

    def walk(content, resources, ctm: tuple[float, ...], depth: int) -> None:
        xobjects = _resolved(_resolved(resources).get("/XObject")) if resources is not None else None
        saved: list[tuple[float, ...]] = []
        for operands, operator in content.operations:
            if operator == b"q":
                saved.append(ctm)
            elif operator == b"Q":
                ctm = saved.pop() if saved else ctm
            elif operator == b"cm" and len(operands) == 6:
                ctm = _multiply(tuple(float(value) for value in operands), ctm)
            elif operator == b"Do" and xobjects is not None:
                xobject = _resolved(xobjects.get(operands[0]))
                subtype = xobject.get("/Subtype") if xobject is not None else None
                if subtype == "/Image":
                    add(_image_key(xobject, keys), ctm)
                elif subtype == "/Form" and depth < 4:
                    matrix = tuple(float(value) for value in _resolved(xobject.get("/Matrix", [1, 0, 0, 1, 0, 0])))
                    inner = ContentStream(xobject, page.pdf)
                    walk(inner, xobject.get("/Resources", resources), _multiply(matrix, ctm), depth + 1)
            elif operator == b"INLINE IMAGE":
                data = operands.get("data", b"") if isinstance(operands, dict) else b""
                add(("inline", hashlib.blake2b(data, digest_size=16).hexdigest()), ctm)

    try:
        content = page.get_contents()
        if content is not None:
            walk(content, page.get("/Resources"), (1.0, 0.0, 0.0, 1.0, 0.0, 0.0), 0)
    except Exception:
        # Drawing instructions that can't be followed: count the page's images, each as possible content.
        return [Picture(key=("unsized", number, index), area=1.0, ratio=1.0) for index in range(_pdf_image_count(page))]
    return found


def _image_key(xobject, keys: dict[int, object]) -> object:
    """The same picture's key wherever it's drawn: by its stored bytes, so a template logo copied
    onto every page is recognised too (hashed once per picture, never decoded)."""
    if id(xobject) not in keys:
        try:
            keys[id(xobject)] = ("image", hashlib.blake2b(xobject._data, digest_size=16).hexdigest())
        except Exception:
            keys[id(xobject)] = ("image", id(xobject))
    return keys[id(xobject)]


def _multiply(m: tuple[float, ...], n: tuple[float, ...]) -> tuple[float, ...]:
    """Two PDF transformation matrices [a b c d e f] combined (m applied first, then n)."""
    a1, b1, c1, d1, e1, f1 = m
    a2, b2, c2, d2, e2, f2 = n
    return (
        a1 * a2 + b1 * c2,
        a1 * b2 + b1 * d2,
        c1 * a2 + d1 * c2,
        c1 * b2 + d1 * d2,
        e1 * a2 + f1 * c2 + e2,
        e1 * b2 + f1 * d2 + f2,
    )


def _resolved(value):
    return value.get_object() if hasattr(value, "get_object") else value


def _pdf_image_count(page) -> int:
    """Images in a PDF page's resources (nothing is decoded)."""
    try:
        resources = page.get("/Resources")
        xobjects = resources.get_object().get("/XObject") if resources else None
        if not xobjects:
            return 0
        return sum(1 for item in xobjects.get_object().values() if item.get_object().get("/Subtype") == "/Image")
    except Exception:
        return 0


def _template_keys(pages: list[list[Picture]]) -> set[object]:
    """Pictures on at least half the pages (and at least three): a slide template's logo or background."""
    counts = Counter(key for pictures in pages for key in {picture.key for picture in pictures})
    needed = max(_TEMPLATE_MIN_PAGES, math.ceil(len(pages) / 2))
    return {key for key, count in counts.items() if count >= needed}


def _content_pictures(pictures: list[Picture], text: str, template: set[object], seen: set[object]) -> list[Picture]:
    """The pictures on a page that may carry content: not the slide template, not an icon or logo.

    A picture shown again on later pages (a slide built up step by step) is read once, where it
    first appears: `seen` collects the pictures already chosen.
    """
    sparse = len(text.split()) <= _SPARSE_PAGE_WORDS
    content = [
        picture
        for picture in pictures
        if picture.key not in template
        and picture.key not in seen
        and (
            picture.area >= _PICTURE_MIN_AREA
            or (picture.ratio >= _STRIP_RATIO and picture.area >= _STRIP_MIN_AREA)
            or (sparse and picture.area >= _SPARSE_MIN_AREA)
        )
    ]
    seen.update(picture.key for picture in content)
    return content


def _worth_reading(number: int, text: str, total: int) -> bool:
    """Whether a page's pictures are worth reading: not on the title slide or a closing one ("Thank you", "Questions?")."""
    words = len(text.split())
    if number == 1 and total > 3 and words <= _TITLE_WORDS:
        return False
    return not (words <= _CLOSING_WORDS and _CLOSING.search(text))


def pdf_subset(data: bytes, pages: list[int]) -> bytes:
    """A PDF of just these pages (1-based), in this order, to show the AI."""
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        reader.decrypt("")
    writer = PdfWriter()
    for number in pages:
        writer.add_page(reader.pages[number - 1])
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def parse_pptx(data: bytes) -> ParsedDocument:
    from pptx import Presentation

    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as error:
        raise DocumentParseError("This PowerPoint file couldn't be opened. It may be damaged.") from error

    units: list[TextUnit] = []
    figures: list[FigureSpot] = []
    slides = list(presentation.slides)
    slide_area = float(presentation.slide_width or 0) * float(presentation.slide_height or 0) or 1.0
    pictures = [list(_pptx_pictures(slide.shapes, slide_area)) for slide in slides]
    template = _template_keys(pictures)
    seen: set[object] = set()
    for number, slide in enumerate(slides, start=1):
        parts: list[str] = []
        title = ""
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = normalize_text(slide.shapes.title.text_frame.text)
        except Exception:
            title = ""
        if title:
            parts.append(title)
        for shape in slide.shapes:
            parts.extend(_shape_texts(shape, skip_text=title))
        charts = [_chart_text(chart, number) for chart in _charts(slide.shapes)]
        parts.extend(chart for chart in charts if chart)
        visible = normalize_text("\n".join(parts))
        try:
            if slide.has_notes_slide:
                notes = normalize_text(slide.notes_slide.notes_text_frame.text)
                if notes:
                    parts.append(f"Speaker notes: {notes}")
        except Exception:
            pass
        worth = _worth_reading(number, visible, len(slides))
        content = _content_pictures(pictures[number - 1], visible, template, seen) if worth else []
        if content:
            figures.append(FigureSpot(number=number, kind="Slide", text=normalize_text("\n".join(parts)), pictures=content))
        # Pictures (and charts whose data couldn't be read) aren't text: say so, so their content is
        # never assumed to be covered (until the pictures are read).
        unread = len(content) + charts.count("")
        if unread and any(parts):
            parts.append(figure_marker("slide", number, unread))
        text = normalize_text("\n".join(part for part in parts if part))
        if text:
            units.append(TextUnit(text=text, page_start=number, page_end=number, kind="Slide"))
    return ParsedDocument(units=units, page_count=len(slides), figures=figures)


FIGURE_MARKER = re.compile(r"\[Figure on (?:slide|page) (\d+):")
# The whole line of a "not in the extracted text" marker, and the number of pictures it counts.
UNREAD_FIGURE_LINE = re.compile(r"^\[Figure on (?:slide|page) \d+: (\d+) picture\(s\) or chart\(s\)[^\]\n]*\]$", re.M)
# A picture ProfPilot has read (see figure_reading.py): this header, what the picture shows, then
# FIGURE_END. A page whose pictures are only decoration has the header alone, of kind DECORATIVE.
READ_FIGURE = re.compile(r"^\[Figure on (page|slide) (\d+), read by ProfPilot: ([^\]\n]+)\]$", re.M)
FIGURE_END = "[End of figure]"
DECORATIVE = "decorative only"
# Unicode "mathematical alphanumeric" letters (𝒇, 𝒘, 𝟏…): what PDF text layers yield for typeset
# maths, often with garbled signs around them. Where they appear, extracted formulas are unreliable.
EXTRACTION_GLYPHS = re.compile("[\U0001d400-\U0001d7ff]")
# "[Page 14]" / "[Slide 3]": where a page or slide starts inside a chunk (see chunk_units).
PAGE_MARKER = re.compile(r"^\[(Page|Slide) (\d+)\]$", re.M)


def figure_marker(unit: str, number: int, count: int) -> str:
    return f"[Figure on {unit} {number}: {count} picture(s) or chart(s) whose content is not in the extracted text]"


def _pptx_pictures(shapes, slide_area: float) -> Iterator[Picture]:
    """The pictures on a slide (inside groups too), with the share of the slide each covers."""
    for shape in shapes:
        try:
            if shape.shape_type == 6 and hasattr(shape, "shapes"):  # a group of shapes
                yield from _pptx_pictures(shape.shapes, slide_area)
                continue
            image = getattr(shape, "image", None)  # a picture, or a placeholder holding one
            if image is None:
                continue
            width, height = float(shape.width or 0), float(shape.height or 0)
            yield Picture(
                key=image.sha1,
                area=min(1.0, width * height / slide_area),
                ratio=max(width, height) / max(min(width, height), 1.0),
                image=(image.content_type, image.blob),
            )
        except Exception:  # a linked (not embedded) picture, or a shape python-pptx can't describe
            continue


def _charts(shapes) -> Iterator:
    for shape in shapes:
        try:
            if shape.shape_type == 6 and hasattr(shape, "shapes"):  # a group of shapes
                yield from _charts(shape.shapes)
            elif getattr(shape, "has_chart", False) and shape.has_chart:
                yield shape.chart
        except Exception:
            continue


_CHART_CATEGORIES = 60
_CHART_SERIES = 12


def _chart_text(chart, number: int) -> str:
    """A PowerPoint chart's own data as text (kind, title, categories, each series' values), or ""."""
    try:
        kind = (getattr(chart.chart_type, "name", None) or "chart").replace("_", " ").lower()
        title = chart.chart_title.text_frame.text.strip() if chart.has_title and chart.chart_title.has_text_frame else ""
        lines = [f"[Chart on slide {number} (data from the file): {kind}{f': {title}' if title else ''}]"]
        for plot in chart.plots:
            categories = [str(category) for category in list(plot.categories)[:_CHART_CATEGORIES]]
            if categories:
                lines.append("Categories: " + " | ".join(categories))
            for series in list(plot.series)[:_CHART_SERIES]:
                values = ["" if value is None else f"{value:g}" for value in list(series.values)[:_CHART_CATEGORIES]]
                lines.append(f"{series.name or 'Series'}: " + " | ".join(values))
        return "\n".join(lines) if len(lines) > 1 else ""
    except Exception:
        return ""


def _shape_texts(shape, *, skip_text: str = "") -> list[str]:
    texts: list[str] = []
    try:
        if shape.shape_type == 6 and hasattr(shape, "shapes"):  # a group of shapes
            for child in shape.shapes:
                texts.extend(_shape_texts(child, skip_text=skip_text))
            return texts
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text = normalize_text(shape.text_frame.text)
            if text and text != skip_text:
                texts.append(text)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [normalize_text(cell.text) for cell in row.cells]
                if any(cells):
                    texts.append(" | ".join(cells))
    except Exception:
        pass
    return texts


def parse_docx(data: bytes) -> ParsedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as error:
        raise DocumentParseError("This Word file couldn't be opened. It may be damaged.") from error

    units: list[TextUnit] = []
    heading: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        text = normalize_text("\n".join(buffer))
        if text:
            units.append(TextUnit(text=text, kind="Section", heading=heading))
        buffer.clear()

    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text.strip()
            style = (block.style.name if block.style is not None else "") or ""
            if style.lower().startswith(("heading", "title")) and text:
                flush()
                heading = text[:200]
                buffer.append(text)
            elif text:
                buffer.append(text)
        elif isinstance(block, Table):
            for row in block.rows:
                cells = [normalize_text(cell.text) for cell in row.cells]
                # Merged cells repeat; keep each value once per row.
                unique = list(dict.fromkeys(cell for cell in cells if cell))
                if unique:
                    buffer.append(" | ".join(unique))
    flush()
    return ParsedDocument(units=units)


def parse_txt(data: bytes) -> ParsedDocument:
    text = None
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            candidate = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        if encoding == "utf-16" and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
            continue
        text = candidate
        break
    if text is None:
        raise DocumentParseError("This text file couldn't be read.")
    text = normalize_text(text)
    units = [TextUnit(text=part) for part in re.split(r"\n{2,}", text) if part.strip()]
    return ParsedDocument(units=units)


# -----------------------------------------------------------------------------
# Chunking
# -----------------------------------------------------------------------------


@dataclass
class Chunk:
    text: str
    page_start: int | None
    page_end: int | None
    location_label: str | None


def location_label(kind: str | None, start: int | None, end: int | None, heading: str | None = None) -> str | None:
    if kind in ("Slide", "Page") and start is not None:
        if end is None or end == start:
            return f"{kind} {start}"
        return f"{kind}s {start}–{end}"
    if heading:
        return f"Section: {heading}"
    return None


def chunk_units(units: list[TextUnit], *, target_chars: int = 2400, overlap_chars: int = 300) -> list[Chunk]:
    """Groups units into retrieval-sized chunks without mixing far-apart text.

    Small consecutive units (e.g. short slides) are combined; long units are
    split at paragraph or sentence boundaries with a little overlap so no idea
    is cut in half without context. Each page or slide starts with a marker
    ("[Page 14]"), so text inside a multi-page chunk keeps its own page.
    """
    chunks: list[Chunk] = []
    current: list[TextUnit] = []
    size = 0

    def flush() -> None:
        nonlocal current, size
        if not current:
            return
        text = "\n\n".join(_marked(unit, unit.text) for unit in current)
        first, last = current[0], current[-1]
        kind = first.kind if all(unit.kind == first.kind for unit in current) else None
        start = first.page_start
        end = last.page_end if last.page_end is not None else start
        chunks.append(
            Chunk(text=text, page_start=start, page_end=end, location_label=location_label(kind, start, end, first.heading))
        )
        current, size = [], 0

    for unit in units:
        text = unit.text
        if len(text) > target_chars:
            flush()
            for piece in _split_long(text, target_chars, overlap_chars):
                chunks.append(
                    Chunk(
                        text=_marked(unit, piece),
                        page_start=unit.page_start,
                        page_end=unit.page_end,
                        location_label=location_label(unit.kind, unit.page_start, unit.page_end, unit.heading),
                    )
                )
            continue
        # Don't merge Word sections with different headings: they are separate topics. Don't
        # merge across a gap (a page with no text) either, so "Pages 3–5" never claims a missing page.
        previous = current[-1] if current else None
        gap = bool(
            previous
            and unit.page_start is not None
            and previous.page_end is not None
            and unit.page_start != previous.page_end + 1
        )
        if current and (size + len(text) > target_chars or unit.heading != current[-1].heading or gap):
            flush()
        current.append(unit)
        size += len(text) + 2
    flush()
    return chunks


def _marked(unit: TextUnit, text: str) -> str:
    """The text of a page or slide, starting with its marker ("[Slide 3]")."""
    if unit.kind in ("Page", "Slide") and unit.page_start is not None:
        return f"[{unit.kind} {unit.page_start}]\n{text}"
    return text


def _split_long(text: str, target: int, overlap: int) -> list[str]:
    paragraphs = [part for part in re.split(r"\n{2,}|\n", text) if part.strip()]
    pieces: list[str] = []
    buffer = ""
    for paragraph in paragraphs:
        for sentence in _sentences(paragraph, target):
            if buffer and len(buffer) + len(sentence) + 1 > target:
                pieces.append(buffer.strip())
                buffer = buffer[-overlap:] if overlap else ""
            buffer = f"{buffer}\n{sentence}" if buffer else sentence
    if buffer.strip():
        pieces.append(buffer.strip())
    return pieces


def _sentences(paragraph: str, target: int) -> list[str]:
    if len(paragraph) <= target:
        return [paragraph]
    sentences = re.split(r"(?<=[.!?])\s+", paragraph)
    out: list[str] = []
    for sentence in sentences:
        while len(sentence) > target:  # a very long "sentence" (e.g. a data dump)
            out.append(sentence[:target])
            sentence = sentence[target:]
        if sentence:
            out.append(sentence)
    return out


def estimate_tokens(text: str) -> int:
    return max(1, len(text) // 4)
