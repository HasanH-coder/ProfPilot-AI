"""Text extraction keeps page, slide and section locations; bad files fail clearly."""

import io

import docx
import pytest
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter

from app.services.document_parsers import (
    DOCX,
    PDF,
    PPTX,
    TXT,
    DocumentParseError,
    TextUnit,
    chunk_units,
    parse_document,
)
from tests.seed import pdf_bytes


def test_pdf_pages_keep_their_numbers():
    parsed = parse_document(pdf_bytes(["Gradient descent basics and learning rates", "Regularization with an L2 penalty"]), PDF)
    assert parsed.page_count == 2
    assert [(unit.page_start, unit.kind) for unit in parsed.units] == [(1, "Page"), (2, "Page")]
    assert "learning rates" in parsed.units[0].text


def test_scanned_pdf_pages_are_flagged_for_transcription():
    parsed = parse_document(pdf_bytes(["Real text on the first page here", ""]), PDF)
    assert parsed.pages_without_text == [2]


def test_password_protected_pdf_fails_with_a_clear_message():
    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(pdf_bytes(["secret"]))))
    writer.encrypt("hunter2")
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(DocumentParseError, match="password"):
        parse_document(buffer.getvalue(), PDF)


def test_damaged_files_fail_with_a_clear_message():
    with pytest.raises(DocumentParseError):
        parse_document(b"%PDF-1.4 this is not really a pdf", PDF)
    with pytest.raises(DocumentParseError):
        parse_document(b"not a zip", PPTX)
    with pytest.raises(DocumentParseError):
        parse_document(b"not a zip", DOCX)


def test_pptx_slides_titles_tables_and_notes():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Backpropagation"
    slide.placeholders[1].text = "Chain rule applied layer by layer"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Layer", "Gradient"
    slide.notes_slide.notes_text_frame.text = "Mention vanishing gradients"
    buffer = io.BytesIO()
    presentation.save(buffer)
    parsed = parse_document(buffer.getvalue(), PPTX)
    assert parsed.page_count == 1
    unit = parsed.units[0]
    assert unit.kind == "Slide" and unit.page_start == 1
    for expected in ("Backpropagation", "Chain rule", "Layer | Gradient", "Speaker notes: Mention vanishing gradients"):
        assert expected in unit.text


def test_docx_sections_follow_headings():
    document = docx.Document()
    document.add_heading("Week 1: Probability", level=1)
    document.add_paragraph("Bayes' rule relates conditional probabilities.")
    document.add_heading("Week 2: Estimation", level=1)
    document.add_paragraph("Maximum likelihood chooses the most likely parameters.")
    buffer = io.BytesIO()
    document.save(buffer)
    parsed = parse_document(buffer.getvalue(), DOCX)
    assert [unit.heading for unit in parsed.units] == ["Week 1: Probability", "Week 2: Estimation"]
    chunks = chunk_units(parsed.units)
    assert chunks[0].location_label == "Section: Week 1: Probability"


def test_text_files_in_common_encodings():
    assert "café" in parse_document("Notes about café data".encode("cp1252"), TXT).units[0].text
    assert "Σ" in parse_document("Sum Σ".encode("utf-8-sig"), TXT).units[0].text


def test_chunks_combine_short_slides_and_split_long_pages():
    slides = [TextUnit(text=f"Slide {n} text " * 10, page_start=n, page_end=n, kind="Slide") for n in range(1, 5)]
    chunks = chunk_units(slides, target_chars=400)
    # Three 130-character slides fit in 400 characters; the fourth starts a new chunk.
    assert [chunk.location_label for chunk in chunks] == ["Slides 1–3", "Slide 4"]
    long_page = TextUnit(text=("A sentence about optimization. " * 200).strip(), page_start=7, page_end=7, kind="Page")
    pieces = chunk_units([long_page], target_chars=1000, overlap_chars=100)
    assert len(pieces) > 3
    assert all(len(piece.text) <= 1100 for piece in pieces)
    assert all(piece.location_label == "Page 7" for piece in pieces)
